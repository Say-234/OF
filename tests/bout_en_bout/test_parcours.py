"""J9-J10 — parcours A (le collaborateur prépare) et B (l'expert valide), à travers les cas d'usage (doc 04).

On joue le parcours complet sur la démo BATI-PLUS avec une vraie base SQLite temporaire,
une horloge figée au 23/09/2026 et les vrais générateurs de livrables.
"""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from conftest import DEMO, M
from optifisc.composition import construire
from optifisc.domaine.dossier import Ajustement, CategorieTaux, EtatDossier, Hypotheses, Secteur
from optifisc.domaine.erreurs import AccesRefuse, DossierFige, EtapeIncomplete
from optifisc.domaine.scenario import LevierChoisi
from optifisc.domaine.utilisateur import Role


class HorlogeFigee:
    def maintenant(self):
        return datetime(2026, 9, 23, 10, 0)

    def aujourd_hui(self):
        return date(2026, 9, 23)


@pytest.fixture
def app(tmp_path):
    return construire(tmp_path / "parcours.sqlite", horloge=HorlogeFigee())


@pytest.fixture
def equipe(app):
    admin = app.comptes.connecter("admin@cabinet.bj", app.comptes.creer_premier_administrateur("admin@cabinet.bj", "A"))
    mdp_salim = app.comptes.creer_utilisateur(admin, "salim@cabinet.bj", "Salim A.", Role.COLLABORATEUR)
    mdp_expert = app.comptes.creer_utilisateur(admin, "expert@cabinet.bj", "Expert", Role.EXPERT)
    mdp_autre = app.comptes.creer_utilisateur(admin, "autre@cabinet.bj", "Autre", Role.COLLABORATEUR)
    return {"admin": admin, "salim": app.comptes.connecter("salim@cabinet.bj", mdp_salim),
            "expert": app.comptes.connecter("expert@cabinet.bj", mdp_expert),
            "autre": app.comptes.connecter("autre@cabinet.bj", mdp_autre)}


def preparer_batiplus(app, salim):
    """Parcours A jusqu'au calcul et à la décision sur les acomptes."""
    d = app.dossiers.creer(salim, "BATI-PLUS SARL", Secteur.BTP, CategorieTaux.DROIT_COMMUN, 2026, demonstration=True)
    contenu = (DEMO / "balance_BATIPLUS_BTP_30-06-2026.xlsx").read_bytes()
    apercu = app.balance.apercu(salim, d.id, contenu, "balance_BATIPLUS.xlsx", date(2026, 6, 30))
    app.balance.confirmer(salim, d.id, apercu)
    app.dossiers.definir_hypotheses(salim, d.id, Hypotheses(
        mois_arrete=6, is_n_moins_1=Decimal(52 * M), aib_imputable=Decimal(4 * M),
        ajustements=(Ajustement("6813", "Dotations 2026 à passer", Decimal(45 * M)),)))
    for ligne in app.calcul.resultat(salim, d.id).retraitements:
        if not ligne.renseignee:
            app.dossiers.saisir(salim, d.id, ligne.id_regle, Decimal(0))       # « néant »
    resultat = app.calcul.calculer(salim, d.id)
    app.calcul.decider_acomptes(salim, d.id, demander=True)
    return d.id, resultat


def test_parcours_complet_collaborateur_puis_expert(app, equipe):
    salim, expert = equipe["salim"], equipe["expert"]
    dossier_id, resultat = preparer_batiplus(app, salim)
    assert resultat.is_du == Decimal("36.84") * M                       # identique à la démo (jalon de vérité)

    # leviers : au minimum de perception, une provision ne change pas l'impôt
    choix = [LevierChoisi("L03_PROVISION_JUSTIFIEE", {"montant": "5000000", "justifiee": True, "au_releve": True})]
    assert app.scenarios.simuler(salim, dossier_id, choix).gain_is == 0
    scenario = app.scenarios.enregistrer(salim, dossier_id, choix)
    app.scenarios.retenir(salim, dossier_id, scenario.id)

    # livrables provisoires pour le collaborateur
    provisoires = app.livrables.generer(salim, dossier_id, definitif=False)
    assert any(lv.nom.endswith("_PROJET.docx") for lv in provisoires)
    with pytest.raises(AccesRefuse):
        app.livrables.generer(salim, dossier_id, definitif=True)

    # soumission → le dossier est figé pour le collaborateur
    app.validation.soumettre(salim, dossier_id)
    with pytest.raises(DossierFige):
        app.dossiers.saisir(salim, dossier_id, "R06_TVM", Decimal(1))

    # validation par l'expert (règles « à valider » utilisées : case à cocher obligatoire)
    with pytest.raises(EtapeIncomplete, match="non validée"):
        app.validation.valider(expert, dossier_id)
    app.validation.valider(expert, dossier_id, regles_non_validees_verifiees=True)
    definitifs = app.livrables.generer(expert, dossier_id, definitif=True)
    assert app.dossiers.ouvrir(expert, dossier_id).etat is EtatDossier.LIVRE
    noms = [lv.nom for lv in definitifs]
    assert any("Lettre_DGI" in n for n in noms)                         # demande de dispense décidée

    # contenu des livrables : le tableau de passage Excel et la note Word portent l'IS dû
    excel = next(lv for lv in definitifs if lv.nom.endswith(".xlsx"))
    classeur = load_workbook(BytesIO(excel.contenu))
    assert {"Synthèse", "Tableau de passage", "Acomptes", "Balance projetée", "Hypothèses", "Piste d'audit"} \
        <= set(classeur.sheetnames)
    note = next(lv for lv in definitifs if lv.nom.startswith("Note"))
    texte = "\n".join(p.text for p in Document(BytesIO(note.contenu)).paragraphs)
    assert "36 840 000 FCFA" in texte and "validée par expert@cabinet.bj" in texte

    # piste d'audit : chaque étape du parcours est tracée
    actions = {e.action for e in app.dependances.journal.lister(dossier_id, 100)}
    assert {"dossier_cree", "balance_importee", "calcul_effectue", "decision_acomptes", "dossier_soumis",
            "dossier_valide", "livrables_definitifs"} <= actions


def test_soumission_refusee_si_saisies_manquantes(app, equipe):
    salim = equipe["salim"]
    d = app.dossiers.creer(salim, "NOVA", Secteur.AUTRES, CategorieTaux.DROIT_COMMUN, 2026)
    contenu = (DEMO / "balance_NOVASERVICES_30-06-2026.xlsx").read_bytes()
    app.balance.confirmer(salim, d.id, app.balance.apercu(salim, d.id, contenu, "n.xlsx", date(2026, 6, 30)))
    app.calcul.calculer(salim, d.id)
    with pytest.raises(EtapeIncomplete, match="non renseignée"):
        app.validation.soumettre(salim, d.id)


def test_cloisonnement_entre_collaborateurs(app, equipe):
    dossier_id, _ = preparer_batiplus(app, equipe["salim"])
    with pytest.raises(AccesRefuse):
        app.dossiers.ouvrir(equipe["autre"], dossier_id)
    assert [r.dossier.id for r in app.dossiers.lister(equipe["autre"])] == []
    assert [r.dossier.id for r in app.dossiers.lister(equipe["expert"])] == [dossier_id]


def test_modifier_une_donnee_rend_le_calcul_perime(app, equipe):
    salim = equipe["salim"]
    dossier_id, _ = preparer_batiplus(app, salim)
    assert app.dossiers.ouvrir(salim, dossier_id).etat is EtatDossier.CALCULE
    app.dossiers.saisir(salim, dossier_id, "R06_TVM", Decimal(1 * M))
    assert app.dossiers.ouvrir(salim, dossier_id).etat is EtatDossier.EN_PREPARATION


def test_validation_et_caducite_d_une_regle(app, equipe):
    expert = equipe["expert"]
    app.validation.valider_regle(expert, "R02_AMENDES_PENALITES", "Amendes 2,4 M → réintégration 2,4 M")
    assert "R02_AMENDES_PENALITES" in app.validation.regles_validees()
    with pytest.raises(AccesRefuse):
        app.validation.valider_regle(equipe["salim"], "R03_DONS_LIBERALITES", "cas")


def test_etapes_du_parcours(app, equipe):
    salim = equipe["salim"]
    dossier_id, _ = preparer_batiplus(app, salim)
    etapes = {e.id: e for e in app.dossiers.etapes(app.dossiers.ouvrir(salim, dossier_id))}
    assert etapes["balance"].faite and etapes["hypotheses"].faite and etapes["calcul"].faite
    assert etapes["acomptes"].faite and not etapes["leviers"].faite and not etapes["livrer"].faite


def test_raccourcis_de_demonstration_refuses_en_cabinet(app, equipe):
    with pytest.raises(AccesRefuse, match="démonstration"):
        app.entrer_en_demo(Role.EXPERT)                    # le parcours fixture tourne en mode « cabinet »


def test_tout_a_neant_repond_aux_saisies_restantes(app, equipe):
    salim = equipe["salim"]
    d = app.dossiers.creer(salim, "TEST", Secteur.AUTRES, CategorieTaux.DROIT_COMMUN, 2026)
    contenu = (DEMO / "balance_NOVASERVICES_30-06-2026.xlsx").read_bytes()
    app.balance.confirmer(salim, d.id, app.balance.apercu(salim, d.id, contenu, "n.xlsx", date(2026, 6, 30)))
    assert app.dossiers.saisir_neant_restantes(salim, d.id) == 23
    assert app.calcul.resultat(salim, d.id).saisies_manquantes == ()
