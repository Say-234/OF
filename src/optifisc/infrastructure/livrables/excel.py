"""Dossier de calcul Excel (repris de la V2) : synthèse, tableau de passage avec totaux EN FORMULES,
acomptes, balance projetée, hypothèses, piste d'audit. Les formules rendent le calcul vérifiable dans Excel."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from optifisc.application.livrables import DonneesLivrables
from optifisc.domaine.montant import en_pourcentage, en_texte
from optifisc.domaine.moteur.projection import facteur_projection

MARINE, BLEU_CLAIR, GRIS = "0F172A", "E8EEF6", "F1F5F9"
NB = '#,##0;[Red]-#,##0;"-"'          # 1 234 567 / -1 234 / -
PCT = "0.0%"
DATE = "DD/MM/YYYY"
MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
        "novembre", "décembre")
AVERTISSEMENT = ("Document d'aide à la décision produit par OptiFisc (prototype). Les montants résultent d'une "
                 "projection et de règles paramétrées à partir du CGI 2026 ; ils doivent être validés par "
                 "l'expert-comptable avant toute communication au client ou démarche auprès de l'administration.")


class _Feuille:
    """Écriture ligne à ligne avec les styles du cabinet."""

    def __init__(self, ws, largeurs: list[int], provisoire: bool):
        self.ws, self.ligne_courante = ws, 1
        for i, largeur in enumerate(largeurs, 1):
            ws.column_dimensions[get_column_letter(i)].width = largeur
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = "OptiFisc — aide à la décision, à valider par l'expert-comptable"
        ws.oddFooter.right.text = "Page &P / &N"
        if provisoire:
            ws.oddHeader.center.text = "PROJET — non validé par l'expert-comptable"

    def titre(self, texte: str, sous_titre: str = ""):
        self.ws.cell(self.ligne_courante, 1, texte).font = Font(size=14, bold=True, color=MARINE)
        self.ligne_courante += 1
        if sous_titre:
            self.ws.cell(self.ligne_courante, 1, sous_titre).font = Font(size=10, italic=True, color="475569")
            self.ligne_courante += 1
        self.ligne_courante += 1

    def entete(self, valeurs: list):
        for j, v in enumerate(valeurs, 1):
            c = self.ws.cell(self.ligne_courante, j, v)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=MARINE)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        self.ws.freeze_panes = self.ws.cell(self.ligne_courante + 1, 1)
        self.ligne_courante += 1

    def section(self, texte: str, n_col: int):
        for j in range(1, n_col + 1):
            self.ws.cell(self.ligne_courante, j).fill = PatternFill("solid", fgColor=BLEU_CLAIR)
        self.ws.cell(self.ligne_courante, 1, texte).font = Font(bold=True, color=MARINE)
        self.ligne_courante += 1

    def ligne(self, valeurs: list, formats: dict | None = None, gras: bool = False, total: bool = False) -> int:
        for j, v in enumerate(valeurs, 1):
            c = self.ws.cell(self.ligne_courante, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            if formats and j in formats:
                c.number_format = formats[j]
            if gras or total:
                c.font = Font(bold=True)
            if total:
                c.border = Border(top=Side(style="thin", color=MARINE))
                c.fill = PatternFill("solid", fgColor=GRIS)
        self.ligne_courante += 1
        return self.ligne_courante - 1

    def texte(self, texte: str, italique: bool = False):
        c = self.ws.cell(self.ligne_courante, 1, texte)
        c.font = Font(italic=italique, color="475569" if italique else MARINE)
        self.ligne_courante += 1


def dossier_de_calcul(d: DonneesLivrables) -> bytes:
    r, dos, p = d.resultat, d.dossier, d.referentiel.parametres
    classeur = Workbook()
    periode = f"situation au {d.date_arrete:%d/%m/%Y}, projetée au 31/12/{dos.exercice}"

    # ---------------------------------------------------------------- tableau de passage (formules)
    ws = classeur.active
    ws.title = "Tableau de passage"
    f = _Feuille(ws, [9, 58, 18, 18, 18, 60], d.provisoire)
    f.titre(f"Détermination de l'impôt sur les sociétés {dos.exercice} — {dos.raison_sociale}", periode)
    f.entete(["Réf.", "Libellé", "Article CGI", "Base", "Montant (FCFA)", "Détail du calcul"])
    ref = {"A": f.ligne(["A", "Résultat comptable projeté au 31/12", "SYSCOHADA", None, r.resultat_comptable,
                         "Produits − charges des classes 6, 7 et 8, projetés sur 12 mois"], {5: NB}, gras=True)}
    categories = {c.id: c.libelle for c in d.referentiel.categories}
    for sens, lettre, titre in (("reintegration", "B", "Réintégrations"), ("deduction", "C", "Déductions")):
        f.section(f"{lettre}. {titre}", 6)
        debut = f.ligne_courante
        categorie_courante = None
        for li in [x for x in r.retraitements if x.sens.value == sens and x.montant > 0]:
            if li.categorie != categorie_courante:
                categorie_courante = li.categorie
                f.ligne(["", categories.get(li.categorie, "")])
                ws.cell(f.ligne_courante - 1, 2).font = Font(italic=True, color="475569")
            f.ligne([li.id_regle.split("_")[0], li.libelle, li.article, li.base, li.montant, li.detail],
                    {4: NB, 5: NB})
        if f.ligne_courante == debut:
            f.ligne(["", "Aucune", "", None, 0], {5: NB})
        ref[lettre] = f.ligne([lettre, f"Total {titre.lower()}", "", None,
                               f"=SUM(E{debut}:E{f.ligne_courante - 1})"], {5: NB}, total=True)
        f.ligne_courante += 1
    ref["D"] = f.ligne(["D", "Résultat fiscal avant imputation des déficits (A + B − C)", "", None,
                        f"=E{ref['A']}+E{ref['B']}-E{ref['C']}"], {5: NB}, gras=True)
    ref["E"] = f.ligne(["E", "Déficits antérieurs imputés", "CGI art. 43", dos.hypotheses.deficits_anterieurs,
                        f"=MIN(MAX(0,E{ref['D']}),D{f.ligne_courante})"], {4: NB, 5: NB})
    ref["F"] = f.ligne(["F", "RÉSULTAT FISCAL (D − E)", "", None, f"=E{ref['D']}-E{ref['E']}"], {5: NB}, total=True)
    f.ligne_courante += 1
    f.section("Calcul de l'impôt", 6)
    ref["G"] = f.ligne(["G", f"IS de droit commun ({en_pourcentage(r.taux_is)} du résultat fiscal positif)",
                        "CGI art. 46", r.taux_is, f"=ROUND(MAX(0,E{ref['F']})*D{f.ligne_courante},0)"],
                       {4: PCT, 5: NB})
    exclus = ", ".join(p.comptes_exclus_encaissables)
    ref["PE"] = f.ligne(["", "Produits encaissables", "CGI art. 47-2", None, r.produits_encaissables,
                         f"Produits des classes 7 et 8 hors comptes {exclus}"], {5: NB})
    ref["H"] = f.ligne(["H", f"Minimum de perception ({en_pourcentage(r.taux_minimum)}, plancher "
                             f"{en_texte(p.plancher_minimum)})", "CGI art. 47", r.taux_minimum,
                        f"=ROUND(MAX(E{ref['PE']}*D{f.ligne_courante},{p.plancher_minimum}),0)"], {4: PCT, 5: NB})
    ref["I"] = f.ligne(["I", "IS DÛ = le plus élevé de G et H", "CGI art. 46-47", None,
                        f"=MAX(E{ref['G']},E{ref['H']})", f"Régime appliqué : {r.regime}"], {5: NB}, total=True)
    ref["J"] = f.ligne(["J", "Redevance ORTB", p.article_ortb, None, r.redevance_ortb,
                        "Payée avec l'acompte du 10 mars"], {5: NB})
    ref["K"] = f.ligne(["K", "TOTAL À ACQUITTER AU TITRE DE L'EXERCICE (I + J)", "", None,
                        f"=E{ref['I']}+E{ref['J']}"], {5: NB}, total=True)
    if r.hors_resultat:
        f.ligne_courante += 1
        f.section("Coûts hors résultat fiscal (à signaler au client)", 6)
        for c in r.hors_resultat:
            f.ligne([c.id.split("_")[0], c.libelle, c.article, c.base, c.montant,
                     "Droit de timbre sur paiements en espèces"], {4: NB, 5: NB})
    f.ligne_courante += 1
    f.texte(f"Seuil de bascule : l'IS de droit commun dépasse le minimum de perception au-delà d'un résultat "
            f"fiscal de {en_texte(r.seuil_bascule)}.", italique=True)
    f.texte("Les totaux sont des formules : toute correction d'une ligne se répercute sur l'impôt.", italique=True)

    # ---------------------------------------------------------------- synthèse (première feuille)
    ws = classeur.create_sheet("Synthèse", 0)
    f = _Feuille(ws, [46, 22, 70], d.provisoire)
    f.titre(f"Pilotage de l'impôt sur les sociétés {dos.exercice}",
            f"Établi le {d.date_edition:%d/%m/%Y} par {d.redacteur}" + (f" — {d.validation}" if d.validation else
                                                                      " — PROJET non validé"))
    f.entete(["Dossier", "", ""])
    for a, b in (("Raison sociale", dos.raison_sociale),
                 ("Secteur (minimum de perception)", p.libelles_secteur[dos.secteur.value]),
                 ("Taux de l'IS", p.libelles_taux[dos.categorie_taux.value]), ("Balance", periode),
                 ("Version du référentiel", r.version_referentiel)):
        f.ligne([a, b])
    f.ligne_courante += 1
    f.entete(["Chiffres clés", "Montant (FCFA)", "Commentaire"])
    tp = "'Tableau de passage'!"
    for libelle, cle, commentaire in (("Résultat comptable projeté", "A", ""), ("Réintégrations", "B", ""),
                                      ("Déductions", "C", ""), ("Résultat fiscal", "F", ""),
                                      (f"IS de droit commun ({en_pourcentage(r.taux_is)})", "G", "Art. 46"),
                                      (f"Minimum de perception ({en_pourcentage(r.taux_minimum)})", "H", "Art. 47"),
                                      ("IS dû", "I", f"Régime : {r.regime}"),
                                      ("Total avec redevance ORTB", "K", "Art. 47-6")):
        f.ligne([libelle, f"={tp}E{ref[cle]}", commentaire], {2: NB}, total=libelle == "IS dû")
    ech = d.echeancier
    f.ligne_courante += 1
    f.entete(["Acomptes et trésorerie", "Montant (FCFA)", "Commentaire"])
    f.ligne(["Acomptes prévus (1/4 de l'IS N-1 chacun)", ech.total_normal, f"4 × {en_texte(ech.acompte_normal)}"],
            {2: NB})
    f.ligne(["Déjà versés à la date d'édition", ech.deja_paye, ""], {2: NB})
    f.ligne(["AIB imputable", ech.aib_imputable, "Art. 133-4"], {2: NB})
    if ech.recommandation and d.demande_reduction:
        f.ligne(["Trésorerie préservée par la demande au DGI", ech.recommandation.tresorerie_preservee,
                 f"Demande à déposer avant le {ech.recommandation.deposer_avant:%d/%m/%Y} (art. 52)"], {2: NB},
                total=True)
    f.ligne([f"Solde à payer (+) ou crédit (−) au {d.date_depot:%d/%m/%Y}, sans demande", ech.solde_sans_demande,
             "IS dû − acomptes − AIB"], {2: NB})
    if ech.solde_avec_demande is not None and d.demande_reduction:
        f.ligne(["Idem, avec la demande", ech.solde_avec_demande, ""], {2: NB})
    if r.alertes:
        f.ligne_courante += 1
        f.entete(["Points d'attention", "", ""])
        for alerte in r.alertes:
            f.ligne([alerte.message])
            ws.merge_cells(start_row=f.ligne_courante - 1, start_column=1, end_row=f.ligne_courante - 1, end_column=3)
            ws.cell(f.ligne_courante - 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[f.ligne_courante - 1].height = 32
    f.ligne_courante += 1
    ws.cell(f.ligne_courante, 1, AVERTISSEMENT).font = Font(italic=True, size=9, color="475569")
    ws.merge_cells(start_row=f.ligne_courante, start_column=1, end_row=f.ligne_courante, end_column=3)
    ws.cell(f.ligne_courante, 1).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[f.ligne_courante].height = 45

    # ---------------------------------------------------------------- acomptes
    ws = classeur.create_sheet("Acomptes")
    f = _Feuille(ws, [16, 20, 14, 28, 22], d.provisoire)
    f.titre(f"Acomptes d'IS {dos.exercice} (art. 51-52)",
            f"Base : IS {dos.exercice - 1} de {en_texte(dos.hypotheses.is_n_moins_1)}")
    f.entete(["Échéance", "Acompte normal", "Statut", "Demande de réduction avant le", "Acompte demandé"])
    debut = f.ligne_courante
    for e in ech.echeances:
        demande = (ech.recommandation.acompte_demande if ech.recommandation and d.demande_reduction and not e.payee
                   else e.montant_normal)
        f.ligne([e.date, e.montant_normal, "payé" if e.payee else "à venir", e.date_limite_demande, demande],
                {1: DATE, 2: NB, 4: DATE, 5: NB})
    f.ligne(["Total", f"=SUM(B{debut}:B{f.ligne_courante - 1})", "", "",
             f"=SUM(E{debut}:E{f.ligne_courante - 1})"], {2: NB, 5: NB}, total=True)

    # ---------------------------------------------------------------- balance projetée
    ws = classeur.create_sheet("Balance projetée")
    f = _Feuille(ws, [12, 46, 11, 20, 20, 26], d.provisoire)
    facteur = facteur_projection(dos.hypotheses)
    f.titre("Comptes de gestion projetés au 31/12", f"Facteur de projection : × {facteur:.4f} ; montants positifs.")
    f.entete(["Compte", "Intitulé", "Nature", f"Réalisé au {d.date_arrete:%d/%m}", "Projeté au 31/12", "Méthode"])
    # projeter() garde l'ordre des comptes de gestion puis ajoute les ajustements à la fin :
    # la i-ème ligne projetée correspond donc à la i-ème ligne réalisée.
    realisees = d.balance_initiale.filtrer(p.plan.comptes_de_gestion)
    lignes = []
    for i, li in enumerate(r.balance_projetee.lignes):
        if i < len(realisees):
            non_extrapole = li.commence_par(p.plan.impot_sur_resultat)
            methode = "non extrapolé (IS)" if non_extrapole else f"réalisé × {facteur:.4f}"
            lignes.append((li, realisees[i].solde, methode))
        else:
            lignes.append((li, 0, "ajustement de fin d'année"))
    debut = f.ligne_courante
    for li, solde_realise, methode in sorted(lignes, key=lambda x: x[0].compte):
        signe = 1 if li.commence_par(p.plan.charges) else -1
        f.ligne([li.compte, li.libelle, "Charge" if signe == 1 else "Produit", signe * solde_realise, signe * li.solde,
                 methode], {4: NB, 5: NB})
    fin = f.ligne_courante - 1
    for nature in ("Produit", "Charge"):
        f.ligne(["", f"Total des {nature.lower()}s", "", f'=SUMIF(C{debut}:C{fin},"{nature}",D{debut}:D{fin})',
                 f'=SUMIF(C{debut}:C{fin},"{nature}",E{debut}:E{fin})'], {4: NB, 5: NB}, total=True)
    f.ligne(["", "Résultat comptable (produits − charges)", "",
             f"=D{f.ligne_courante - 2}-D{f.ligne_courante - 1}", f"=E{f.ligne_courante - 2}-E{f.ligne_courante - 1}"],
            {4: NB, 5: NB}, total=True)

    # ---------------------------------------------------------------- hypothèses
    ws = classeur.create_sheet("Hypothèses")
    f = _Feuille(ws, [52, 30, 60], d.provisoire)
    f.titre("Hypothèses et informations saisies", "À relire : elles conditionnent tous les montants.")
    f.entete(["Élément", "Valeur", "Source / commentaire"])
    h = dos.hypotheses
    f.ligne(["Balance arrêtée à fin", MOIS[h.mois_arrete - 1], d.rapport_lecture.fichier])
    f.ligne(["Rythme d'activité du reste de l'année", float(h.coef_activite), "1 = même rythme mensuel"])
    f.ligne([f"IS dû au titre de {dos.exercice - 1}", h.is_n_moins_1, "Base des acomptes (art. 51)"], {2: NB})
    f.ligne(["AIB supporté et imputable", h.aib_imputable, "Art. 133-4"], {2: NB})
    f.ligne(["Déficits antérieurs reportables", h.deficits_anterieurs, "Art. 43"], {2: NB})
    for a in h.ajustements:
        f.ligne([f"Ajustement de fin d'année : {a.libelle}", a.montant, f"Compte {a.compte}"], {2: NB})
    libelles = {x.id: (x.libelle, x.article) for x in (*d.referentiel.regles, *d.referentiel.sanctions)}
    if dos.saisies:
        f.ligne_courante += 1
        f.section("Montants saisis par le collaborateur (0 = néant)", 3)
        for cle, montant in sorted(dos.saisies.items()):
            libelle, article = libelles.get(cle, (cle, ""))
            f.ligne([libelle, montant, article], {2: NB})
    f.ligne_courante += 1
    f.section("Lecture de la balance", 3)
    lec = d.rapport_lecture
    f.ligne(["Fichier", lec.fichier, lec.feuille])
    f.ligne(["Montants pris dans", lec.base, ""])
    f.ligne(["Comptes lus", len(d.balance_initiale.lignes), ""])
    f.ligne(["Comptes de regroupement ignorés", len(lec.regroupements_ignores),
             ", ".join(lec.regroupements_ignores[:15])])
    f.ligne(["Lignes de total ignorées", lec.lignes_totaux_ignorees, ""])
    if d.regles_non_validees:
        f.ligne_courante += 1
        f.texte(f"Règles utilisées encore « à valider » par le cabinet : {', '.join(d.regles_non_validees)}.",
                italique=True)

    # ---------------------------------------------------------------- piste d'audit
    ws = classeur.create_sheet("Piste d'audit")
    f = _Feuille(ws, [24, 26, 50, 18, 18, 18, 70], d.provisoire)
    f.titre("Piste d'audit du calcul", "Chaque montant avec sa règle, son article et sa formule.")
    f.entete(["Étape", "Référence", "Libellé", "Article", "Base", "Montant", "Détail du calcul"])
    for a in r.audit:
        f.ligne([a.etape, a.reference, a.libelle, a.article, a.base, a.montant, a.detail], {5: NB, 6: NB})

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
