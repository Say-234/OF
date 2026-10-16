"""Note de pilotage et lettre au DGI en Word (reprises de la V2), modifiables avant envoi."""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from optifisc.application.livrables import DonneesLivrables
from optifisc.domaine.montant import en_pourcentage, en_texte

from .excel import AVERTISSEMENT, MOIS

MARINE = "0F172A"


def _millions(montant) -> str:
    """36 840 000 → « 36,8 M FCFA » ; sous le million, montant complet."""
    if abs(montant) >= 1_000_000:
        return f"{montant / 1_000_000:,.1f}".replace(",", " ").replace(".", ",") + " M FCFA"
    return en_texte(montant)


def _fond(cellule, couleur: str):
    proprietes = cellule._tc.get_or_add_tcPr()
    ombrage = OxmlElement("w:shd")
    ombrage.set(qn("w:val"), "clear")
    ombrage.set(qn("w:color"), "auto")
    ombrage.set(qn("w:fill"), couleur)
    proprietes.append(ombrage)


def _document(provisoire: bool) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin = section.bottom_margin = Cm(2)
        if provisoire:                                  # « filigrane » : mention en en-tête de chaque page
            entete = section.header.paragraphs[0]
            run = entete.add_run("PROJET — non validé par l'expert-comptable")
            run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0xB4, 0x53, 0x09)
            entete.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def _tableau(doc, entetes, lignes, largeurs, droite=frozenset(), gras=frozenset()):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, e in enumerate(entetes):
        c = t.rows[0].cells[j]
        run = c.paragraphs[0].add_run(e)
        run.bold, run.font.color.rgb, run.font.size = True, RGBColor(0xFF, 0xFF, 0xFF), Pt(9.5)
        _fond(c, MARINE)
    for i, ligne in enumerate(lignes):
        cellules = t.add_row().cells
        for j, v in enumerate(ligne):
            p = cellules[j].paragraphs[0]
            run = p.add_run(str(v))
            run.font.size, run.bold = Pt(9.5), i in gras
            if j in droite:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i in gras:
                _fond(cellules[j], "F1F5F9")
    for ligne in t.rows:
        for j, largeur in enumerate(largeurs):
            ligne.cells[j].width = Cm(largeur)
    return t


def note_de_pilotage(d: DonneesLivrables) -> bytes:
    r, dos = d.resultat, d.dossier
    doc = _document(d.provisoire)
    p = doc.add_paragraph()
    run = p.add_run(f"NOTE DE PILOTAGE FISCAL — EXERCICE {dos.exercice}")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0x47, 0x55, 0x69)
    doc.add_heading(dos.raison_sociale, level=0)
    doc.add_paragraph(f"Situation au {d.date_arrete:%d/%m/%Y}, projetée au 31/12/{dos.exercice} · établie le "
                      f"{d.date_edition:%d/%m/%Y} par {d.redacteur}")
    v = doc.add_paragraph()
    run = v.add_run(d.validation or "Projet : note non encore validée par l'expert-comptable.")
    run.bold, run.font.size = True, Pt(9.5)
    run.font.color.rgb = RGBColor(0x15, 0x80, 0x3D) if d.validation else RGBColor(0xB4, 0x53, 0x09)

    numero = 1
    doc.add_heading(f"{numero}. Impôt sur les sociétés prévu", level=1)
    p = doc.add_paragraph(f"Impôt sur les sociétés prévu pour {dos.exercice} : ")
    p.add_run(en_texte(r.is_du)).bold = True
    p.add_run(f" (plus la redevance ORTB de {en_texte(r.redevance_ortb)}).")
    if r.au_minimum:
        doc.add_paragraph(f"C'est le minimum de perception (art. 47) qui s'applique : {en_pourcentage(r.taux_minimum)}"
                          f" des produits encaissables ({_millions(r.produits_encaissables)}) donne "
                          f"{_millions(r.minimum_perception)}, davantage que l'impôt calculé sur le bénéfice "
                          f"({en_pourcentage(r.taux_is)} × {_millions(r.resultat_fiscal)} = "
                          f"{_millions(r.is_droit_commun)}).")
    else:
        doc.add_paragraph(f"C'est l'impôt sur le bénéfice (art. 46) qui s'applique : {en_pourcentage(r.taux_is)} du "
                          f"bénéfice fiscal prévu ({_millions(r.resultat_fiscal)}). Le minimum de perception "
                          f"({_millions(r.minimum_perception)}) est dépassé.")
    _tableau(doc, ["Indicateur", "Montant"], [
        ["Bénéfice fiscal prévu", en_texte(r.resultat_fiscal)],
        [f"IS sur le bénéfice ({en_pourcentage(r.taux_is)})", en_texte(r.is_droit_commun)],
        [f"Minimum de perception ({en_pourcentage(r.taux_minimum)})", en_texte(r.minimum_perception)],
        ["IS dû (le plus élevé des deux)", en_texte(r.is_du)],
        ["Seuil de bascule du bénéfice", en_texte(r.seuil_bascule)],
    ], [11, 5], droite={1}, gras={3})

    if d.recommandations:
        numero += 1
        doc.add_heading(f"{numero}. Recommandations", level=1)
        for reco in d.recommandations:
            doc.add_paragraph(style="List Number").add_run(reco.titre).bold = True
            doc.add_paragraph(reco.texte)
    if d.simulation:
        numero += 1
        s = d.simulation
        doc.add_heading(f"{numero}. Scénario d'optimisation retenu ({d.nom_scenario})", level=1)
        doc.add_paragraph(f"Impôt de référence {_millions(s.reference.is_du)} → impôt simulé "
                          f"{_millions(s.simulation.is_du)} : économie de {_millions(s.gain_is)} pour une dépense "
                          f"engagée de {_millions(s.decaissement)}.")
        for texte in s.lecture:
            doc.add_paragraph(texte, style="List Bullet")
    if d.pistes_acceptees:
        numero += 1
        doc.add_heading(f"{numero}. Points à examiner par le cabinet", level=1)
        doc.add_paragraph("Pistes proposées par l'assistant IA, acceptées par le cabinet pour examen. Elles ne sont "
                          "pas chiffrées : à étudier au regard des articles cités.").runs[0].italic = True
        for piste in d.pistes_acceptees:
            doc.add_paragraph(f"{piste['titre']} ({piste['article']}) : {piste['justification']}", style="List Bullet")
    if r.alertes:
        numero += 1
        doc.add_heading(f"{numero}. Points d'attention", level=1)
        for alerte in r.alertes:
            doc.add_paragraph(alerte.message, style="List Bullet")
    numero += 1
    doc.add_heading(f"{numero}. Acomptes d'impôt (art. 51-52)", level=1)
    ech = d.echeancier
    lignes = []
    for e in ech.echeances:
        demande = (en_texte(ech.recommandation.acompte_demande)
                   if ech.recommandation and d.demande_reduction and not e.payee else "")
        lignes.append([f"{e.date:%d/%m/%Y}", en_texte(e.montant_normal), "payé" if e.payee else "à venir", demande])
    _tableau(doc, ["Échéance", "Acompte normal", "Statut", "Montant demandé"], lignes, [3.5, 4, 3, 4], droite={1, 3})

    numero += 1
    doc.add_heading(f"{numero}. Passage du résultat comptable à l'impôt", level=1)
    lignes, gras = [], set()
    for i, a in enumerate(r.audit):
        lignes.append([a.libelle, a.article, en_texte(a.montant)])
        if a.etape.startswith(("1.", "6.")):
            gras.add(i)
    _tableau(doc, ["Étape", "Article", "Montant"], lignes, [10, 3.5, 3.5], droite={2}, gras=gras)
    if d.regles_non_validees:
        doc.add_paragraph(f"Réserve : règles encore « à valider » par le cabinet utilisées dans ce calcul : "
                          f"{', '.join(d.regles_non_validees)}.").runs[0].italic = True
    doc.add_paragraph()
    q = doc.add_paragraph(AVERTISSEMENT)
    q.runs[0].italic, q.runs[0].font.size = True, Pt(8.5)
    tampon = BytesIO()
    doc.save(tampon)
    return tampon.getvalue()


def lettre_dgi(d: DonneesLivrables) -> bytes | None:
    """Projet de lettre (art. 52), seulement si une demande de réduction a été décidée et reste possible."""
    reco = d.echeancier.recommandation
    if not (reco and d.demande_reduction):
        return None
    dos, ech = d.dossier, d.echeancier
    nature = "dispense" if reco.dispense else "réduction"
    texte = f"""{dos.raison_sociale}
[Adresse] — IFU : [à compléter]

À Monsieur le Directeur Général des Impôts
Cotonou, le {d.date_edition:%d/%m/%Y}

Objet : Demande de {nature} d'acompte d'impôt sur les sociétés — exercice {dos.exercice} (article 52 du CGI)

Monsieur le Directeur Général,

En application de l'article 52 du Code général des impôts, nous avons l'honneur de solliciter la {nature} de l'acompte d'impôt sur les sociétés exigible le {reco.prochaine_echeance:%d/%m/%Y}.

Sur la base de notre situation arrêtée au {d.date_arrete:%d/%m/%Y}, l'impôt sur les sociétés dont nous estimons être redevables au titre de l'exercice {dos.exercice} s'élève à {en_texte(reco.is_estime_declare)}. Les acomptes calculés sur l'impôt de l'exercice {dos.exercice - 1} représentent {en_texte(ech.total_normal)}, dont {en_texte(ech.deja_paye)} déjà acquittés.

[Motifs à développer par le cabinet : évolution de l'activité, des marges ; justificatifs joints.]

Nous attestons être à jour de nos obligations fiscales et joignons : situation comptable à fin {MOIS[dos.hypotheses.mois_arrete - 1]} {dos.exercice}, projection du résultat, quittances des acomptes versés.

Veuillez agréer, Monsieur le Directeur Général, l'expression de notre haute considération.

[Signature du représentant légal]"""
    doc = _document(d.provisoire)
    for bloc in texte.split("\n\n"):
        p = doc.add_paragraph()
        lignes = bloc.split("\n")
        for i, ligne in enumerate(lignes):
            run = p.add_run(ligne)
            if ligne.startswith("Objet"):
                run.bold = True
            if "[" in ligne:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW    # à compléter par le cabinet
            if i < len(lignes) - 1:
                run.add_break()
        if bloc.startswith(("À Monsieur", "[Signature")):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tampon = BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
