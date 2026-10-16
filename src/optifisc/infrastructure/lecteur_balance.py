"""Lecture d'une balance générale exportée d'un logiciel comptable (implémente le port LecteurBalance).

Repris de la V2 (OptiFisc_App/optifisc/balance.py), sans pandas et en Decimal.

Formats reconnus :
  - 4 colonnes : Compte | Intitulé | Débit | Crédit
  - 3 colonnes : Compte | Intitulé | Solde (solde signé, convention inverse corrigée)
  - 6 ou 8 colonnes : soldes à nouveau / mouvements / soldes de fin, en débit et crédit,
    avec un en-tête sur une ou deux lignes (cellules fusionnées)
Sont ignorés : les lignes de titre au-dessus du tableau, les lignes de total et les
comptes de regroupement (« 60 » quand « 601100 » existe), pour ne rien compter deux fois.
"""
import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from optifisc.domaine.balance import Balance, LigneBalance, RapportLecture
from optifisc.domaine.erreurs import BalanceIllisible

ENCODAGES = ("utf-8-sig", "cp1252", "latin-1")
LIGNES_EXPLOREES = 40       # l'en-tête est cherché dans les 40 premières lignes
MOTS_COMPTE = ("compte", "comptes", "n compte", "no compte", "num compte", "numero", "numero de compte",
               "numero compte", "code compte", "account")
MOTS_LIBELLE = ("intitule", "libelle", "designation", "nom du compte", "denomination")
PHASES = [  # l'ordre compte : « solde à nouveau » = ouverture, « solde fin de période » = clôture
    ("ouverture", r"\b(a nouveau|nouveaux?|ouverture|debut|initial|anterieurs?|reports?)\b"),
    ("cloture", r"\b(fin|cloture|final|finaux)\b"),
    ("mouvements", r"\b(mouvements?|mvts?|periode|cumuls?)\b"),
    ("cloture", r"\bsoldes?\b"),
]
BASES = {"cloture": "soldes de fin de période", "simple": "colonnes débit / crédit",
         "mouvements": "mouvements de la période", "signe": "solde signé"}
ZERO = Decimal("0")


# ------------------------------------------------------------------ texte
def _vide(x) -> bool:
    return x is None or (isinstance(x, str) and not x.strip())


def _txt(x) -> str:
    return "" if x is None else " ".join(str(x).split())


def _norm(x) -> str:
    """Minuscules, sans accents ni ponctuation : « N° Compte » → « n compte »."""
    texte = unicodedata.normalize("NFKD", _txt(x)).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texte.lower()).split())


# ------------------------------------------------------------------ lecture brute (grille de cellules)
def _grilles(contenu: bytes, nom: str) -> list[tuple[str, list[list], str]]:
    """Renvoie [(nom de feuille, lignes, encodage)] quel que soit le format du fichier."""
    if nom.endswith((".csv", ".txt")):
        return [("CSV", *_grille_csv(contenu))]
    if nom.endswith(".xls"):
        return _grilles_xls(contenu)
    return _grilles_xlsx(contenu)


def _grille_csv(contenu: bytes) -> tuple[list[list], str]:
    for encodage in ENCODAGES:
        try:
            texte = contenu.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    debut = texte.splitlines()[:LIGNES_EXPLOREES]
    separateur = max(";,\t|", key=lambda s: sum(ligne.count(s) for ligne in debut))
    return [list(r) for r in csv.reader(io.StringIO(texte), delimiter=separateur)], encodage


def _grilles_xlsx(contenu: bytes) -> list[tuple[str, list[list], str]]:
    from openpyxl import load_workbook
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — tout échec de lecture devient un message clair
        raise BalanceIllisible("Le fichier n'est pas un classeur Excel lisible (endommagé, protégé par mot de "
                               "passe ou d'un autre format). Enregistrez-le en .xlsx, ou exportez la balance "
                               "en CSV depuis le logiciel comptable.") from exc
    return [(feuille.title, [list(r) for r in feuille.iter_rows(values_only=True)], "") for feuille in classeur]


def _grilles_xls(contenu: bytes) -> list[tuple[str, list[list], str]]:
    import xlrd
    try:
        classeur = xlrd.open_workbook(file_contents=contenu)
    except Exception as exc:  # noqa: BLE001
        raise BalanceIllisible("Fichier .xls illisible : enregistrez-le en .xlsx ou exportez en CSV.") from exc
    return [(f.name, [f.row_values(i) for i in range(f.nrows)], "") for f in classeur.sheets()]


def _nettoyer(grille: list[list]) -> list[list]:
    """Égalise la largeur des lignes et retire les colonnes entièrement vides."""
    largeur = max((len(r) for r in grille), default=0)
    grille = [list(r) + [None] * (largeur - len(r)) for r in grille]
    utiles = [j for j in range(largeur) if any(not _vide(r[j]) for r in grille)]
    return [[r[j] for j in utiles] for r in grille]


# ------------------------------------------------------------------ en-tête
def _est_mot_compte(c: str) -> bool:
    return c in MOTS_COMPTE or c.startswith(("n compte", "no compte", "numero de compte", "num compte"))


def _classer(libelle: str) -> tuple[str | None, str | None]:
    """(sens, phase) d'une colonne : sens = D / C / None ; phase = ouverture / mouvements / cloture / None."""
    sens = "D" if re.search(r"\bdebit(eur)?s?\b", libelle) else "C" if re.search(r"\bcredit(eur)?s?\b", libelle) else None
    phase = next((p for p, motif in PHASES if re.search(motif, libelle)), None)
    return sens, phase


def _nb_sens(cellules: list[str]) -> int:
    return sum(1 for c in cellules if re.search(r"\b(debit|credit|debiteur|crediteur)s?\b", c))


def _trouver_entete(grille: list[list]) -> int | None:
    for i, ligne in enumerate(grille[:LIGNES_EXPLOREES]):
        if any(_est_mot_compte(_norm(x)) for x in ligne):
            return i
    return None


def _decrire_colonnes(grille: list[list], h: int) -> tuple[dict[int, str], int]:
    """{index de colonne: libellé tel qu'écrit} et nombre de lignes d'en-tête (1 ou 2).

    Gère les en-têtes sur deux lignes : « Soldes fin » au-dessus de « Débit | Crédit ».
    """
    vide = [""] * len(grille[h])
    ligne = [_txt(x) for x in grille[h]]
    suivante = [_txt(x) for x in grille[h + 1]] if h + 1 < len(grille) else vide
    precedente = [_txt(x) for x in grille[h - 1]] if h > 0 else vide
    if _nb_sens([_norm(c) for c in suivante]) >= 2:            # « Compte » en haut, Débit/Crédit dessous
        haut, bas, n = ligne, suivante, 2
    elif _nb_sens([_norm(c) for c in ligne]) >= 2 and any(_classer(_norm(c))[1] for c in precedente if c):
        haut, bas, n = precedente, ligne, 1                   # groupes au-dessus de la ligne « Compte »
    else:
        return dict(enumerate(ligne)), 1
    groupe, libelles = "", {}
    for j, (c, b) in enumerate(zip(haut, bas)):
        if c:
            groupe = c
        elif not b:
            groupe = ""
        libelles[j] = " ".join(x for x in ((c or (groupe if b else "")), b) if x) if c != b else c
    return libelles, n


# ------------------------------------------------------------------ valeurs
def code_compte(v) -> str:
    """Numéro de compte en chiffres ; "" si la cellule n'est pas un compte (titre, total…)."""
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else ""
    s = str(v).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    if "total" in _norm(s):
        return ""
    m = re.match(r"^\d[\d .\-]*", s)
    return re.sub(r"\D", "", m.group(0)) if m else ""


def nombre(v) -> Decimal | None:
    """Montant en Decimal ; None si illisible. Gère 1 234 567,89 / 1.234.567 / (1 234) / 1234- / 12 000 FCFA."""
    if v is None or isinstance(v, bool):
        return ZERO
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return ZERO if v != v else Decimal(repr(v))      # v != v : NaN ; repr donne l'écriture la plus courte
    s = str(v).strip()
    if s.lower() in ("", "-", "—", "nan", "none"):
        return ZERO
    negatif = False
    if s.startswith("(") and s.endswith(")"):
        negatif, s = True, s[1:-1]
    if s.endswith("-"):
        negatif, s = True, s[:-1]
    s = re.sub(r"(?i)f\s*cfa|[\s  '’]", "", s)
    if "," in s and "." in s:                            # le dernier séparateur est le séparateur décimal
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")
    try:
        x = Decimal(s)
    except InvalidOperation:
        return None
    return -x if negatif else x


# ------------------------------------------------------------------ normalisation
def _normaliser(grille: list[list]) -> tuple[list[LigneBalance], dict]:
    grille = _nettoyer(grille)
    h = _trouver_entete(grille)
    if h is None:
        raise BalanceIllisible("colonne « Compte » introuvable dans les 40 premières lignes.")
    libelles, n_entete = _decrire_colonnes(grille, h)
    normes = {j: _norm(c) for j, c in libelles.items()}
    j_compte = next(j for j, c in normes.items() if _est_mot_compte(c))
    j_libelle = next((j for j, c in normes.items() if j != j_compte and any(m in c for m in MOTS_LIBELLE)), None)

    paires: dict[str, dict] = {}
    j_signe = None
    for j, c in normes.items():
        if j in (j_compte, j_libelle) or not c:
            continue
        sens, phase = _classer(c)
        if sens:
            paires.setdefault(phase or "simple", {})[sens] = j
        elif phase == "cloture" and j_signe is None:
            j_signe = j
    completes = {p: v for p, v in paires.items() if "D" in v and "C" in v}
    if "cloture" in completes:
        base = "cloture"
    elif "simple" in completes:
        base = "simple"
    elif "mouvements" in completes:
        base = "mouvements"
    elif j_signe is not None:
        base = "signe"
    else:
        raise BalanceIllisible("il faut des colonnes débit et crédit, ou une colonne solde. Colonnes lues : "
                               + ", ".join(c for c in libelles.values() if c))

    lignes, totaux, illisibles = [], 0, []
    for i, rang in enumerate(grille[h + n_entete:], start=h + n_entete + 1):
        code = code_compte(rang[j_compte])
        if not code:
            if any(nombre(rang[j]) for j in libelles if j not in (j_compte, j_libelle)):
                totaux += 1                                 # ligne de total sans numéro de compte
            continue

        def lire(j):
            x = nombre(rang[j])
            if x is None:
                illisibles.append(f"ligne {i} : « {rang[j]} »")
                return ZERO
            return x

        if base == "signe":
            s = lire(j_signe)
            d, c = max(s, ZERO), max(-s, ZERO)
        elif base == "mouvements":
            d, c = lire(completes["mouvements"]["D"]), lire(completes["mouvements"]["C"])
            if "ouverture" in completes:
                d, c = d + lire(completes["ouverture"]["D"]), c + lire(completes["ouverture"]["C"])
        else:
            d, c = lire(completes[base]["D"]), lire(completes[base]["C"])
        lib = _txt(rang[j_libelle]) if j_libelle is not None else ""
        lignes.append([code, lib, d, c])
    if not lignes:
        raise BalanceIllisible("aucune ligne de compte lue sous l'en-tête.")

    # comptes de regroupement : « 60 » ou « 601 » quand un compte plus détaillé existe
    codes = sorted({ligne[0] for ligne in lignes})
    regroupements = {c for c, suivant in zip(codes, codes[1:]) if suivant.startswith(c)}
    lignes = [ligne for ligne in lignes if ligne[0] not in regroupements]

    signe_inverse = False
    if base == "signe":
        charges = sum((d - c for code, _, d, c in lignes if code.startswith("6")), ZERO)
        produits = sum((d - c for code, _, d, c in lignes if code.startswith("7")), ZERO)
        if produits > 0 and charges < 0:                    # convention inverse : produits positifs
            signe_inverse = True
            lignes = [[code, lib, c, d] for code, lib, d, c in lignes]

    colonnes = {"compte": libelles[j_compte], "libelle": libelles.get(j_libelle, "—")}
    if base == "signe":
        colonnes["solde"] = libelles[j_signe]
    else:
        colonnes.update(debit=libelles[completes[base]["D"]], credit=libelles[completes[base]["C"]])
    infos = dict(ligne_entete=h + 1, base=BASES[base], colonnes=colonnes, lignes_totaux_ignorees=totaux,
                 regroupements_ignores=tuple(sorted(regroupements)), valeurs_illisibles=tuple(illisibles),
                 signe_inverse=signe_inverse)
    return [LigneBalance(code, lib, d, c) for code, lib, d, c in lignes], infos


class LecteurBalanceTableur:
    """Lit un fichier Excel (.xlsx, .xls) ou CSV. Implémente le port LecteurBalance."""

    def lire(self, contenu: bytes, nom_fichier: str) -> tuple[Balance, RapportLecture]:
        nom = nom_fichier.lower()
        erreurs = []
        for feuille, grille, encodage in _grilles(contenu, nom):
            try:
                lignes, infos = _normaliser(grille)
            except BalanceIllisible as exc:
                erreurs.append(f"{feuille} : {exc}")
                continue
            return Balance(tuple(lignes)), RapportLecture(fichier=nom_fichier, feuille=feuille,
                                                          encodage=encodage, **infos)
        raise BalanceIllisible(" ; ".join(erreurs) or "fichier vide.")
